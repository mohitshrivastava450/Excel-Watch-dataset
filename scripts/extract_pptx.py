import sys
import os
import zipfile
import xml.etree.ElementTree as ET

def extract_with_pptx(pptx_path):
    from pptx import Presentation
    prs = Presentation(pptx_path)
    out_lines = []
    for i, slide in enumerate(prs.slides, start=1):
        out_lines.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text and shape.text.strip():
                out_lines.append(shape.text.strip())
        try:
            notes_slide = slide.notes_slide
            if notes_slide and notes_slide.notes_text_frame and notes_slide.notes_text_frame.text.strip():
                out_lines.append("Notes: " + notes_slide.notes_text_frame.text.strip())
        except Exception:
            pass
    return "\n".join(out_lines)

def extract_with_zip(pptx_path):
    out_lines = []
    with zipfile.ZipFile(pptx_path, 'r') as z:
        slide_files = sorted([f for f in z.namelist() if f.startswith('ppt/slides/slide') and f.endswith('.xml')])
        for idx, sf in enumerate(slide_files, start=1):
            out_lines.append(f"--- Slide {idx} ---")
            data = z.read(sf)
            root = ET.fromstring(data)
            # find all text nodes (a:t)
            for t in root.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}t'):
                text = (t.text or '').strip()
                if text:
                    out_lines.append(text)
        # notes
        notes_files = sorted([f for f in z.namelist() if f.startswith('ppt/notesSlides/notesSlide') and f.endswith('.xml')])
        for nf in notes_files:
            data = z.read(nf)
            root = ET.fromstring(data)
            texts = [ (t.text or '').strip() for t in root.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}t') ]
            joined = ' '.join([t for t in texts if t])
            if joined:
                out_lines.append('Notes: ' + joined)
    return "\n".join(out_lines)

def extract(pptx_path):
    try:
        return extract_with_pptx(pptx_path)
    except Exception:
        return extract_with_zip(pptx_path)

def main():
    pptx_path = sys.argv[1] if len(sys.argv) > 1 else os.path.join('Result', 'Flipkartppt.pptx')
    if not os.path.exists(pptx_path):
        print(f"PPTX not found: {pptx_path}")
        sys.exit(2)
    text = extract(pptx_path)
    out_path = os.path.splitext(pptx_path)[0] + '_extracted.txt'
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(text)
    print(out_path)

if __name__ == '__main__':
    main()
